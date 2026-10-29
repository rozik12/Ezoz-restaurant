"""Presentations: the model writes the outline as JSON, python-pptx turns it into a clean 16:9 .pptx."""
from __future__ import annotations

import io
import json
import re
from dataclasses import dataclass, field

NAVY = (0x14, 0x23, 0x3F)
ACCENT = (0x1F, 0x9D, 0x8B)
PAPER = (0xF7, 0xF5, 0xF0)
INK = (0x22, 0x2B, 0x38)
MUTED = (0x6B, 0x74, 0x80)

LANG_NAME = {"ru": "Russian", "uz": "Uzbek (Latin script)", "en": "English"}


@dataclass
class Slide:
    title: str
    bullets: list[str] = field(default_factory=list)
    notes: str = ""


@dataclass
class Deck:
    title: str
    subtitle: str
    slides: list[Slide]


def outline_prompt(lang: str, count: int) -> str:
    return (
        "You design student presentations. Reply with ONE JSON object and nothing else, no code fences:\n"
        '{"title": str, "subtitle": str, "slides": [{"title": str, "bullets": [str], "notes": str}]}\n'
        f"Rules: exactly {count} content slides (the title slide is added separately); 3 to 5 bullets per slide, each at most 14 words, "
        "concrete and informative, no filler; the last slide is a short conclusion. 'notes' is 2-3 sentences the student can say aloud. "
        "Never invent statistics, quotes or bibliographic sources; if a number is not well known, describe the idea without it. "
        f"Write everything in {LANG_NAME.get(lang, 'Russian')} unless the topic text is clearly in another language.")


def parse_deck(raw: str) -> Deck:
    m = re.search(r"\{.*\}", raw, re.S)
    if not m:
        raise ValueError("no JSON")
    data = json.loads(m.group(0))
    slides = []
    for s in data.get("slides", []):
        title = str(s.get("title", "")).strip()
        bullets = [str(b).strip() for b in s.get("bullets", []) if str(b).strip()][:6]
        if title:
            slides.append(Slide(title[:120], bullets, str(s.get("notes", "")).strip()))
    if not slides:
        raise ValueError("no slides")
    return Deck(str(data.get("title", "")).strip()[:150] or slides[0].title, str(data.get("subtitle", "")).strip()[:200], slides)


def build_pptx(deck: Deck) -> bytes:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Emu, Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    rgb = lambda c: RGBColor(*c)

    def rect(slide, x, y, w, h, color):
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(color)
        shape.line.fill.background()
        return shape

    def text(slide, x, y, w, h, value, size, color, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
        box = slide.shapes.add_textbox(x, y, w, h)
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        p = tf.paragraphs[0]
        p.alignment = align
        run = p.add_run()
        run.text = value
        run.font.size, run.font.bold = Pt(size), bold
        run.font.color.rgb = rgb(color)
        run.font.name = "Calibri"
        return box

    # title slide
    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, prs.slide_width, prs.slide_height, NAVY)
    rect(s, Inches(0.9), Inches(3.05), Inches(1.6), Inches(0.08), ACCENT)
    text(s, Inches(0.9), Inches(1.2), Inches(11.5), Inches(1.9), deck.title, 44 if len(deck.title) < 60 else 34, (255, 255, 255),
         bold=True, anchor=MSO_ANCHOR.BOTTOM)
    if deck.subtitle:
        text(s, Inches(0.9), Inches(3.3), Inches(11.5), Inches(1.2), deck.subtitle, 22, (0xC9, 0xD6, 0xE8))

    total = len(deck.slides)
    for i, sl in enumerate(deck.slides, 1):
        s = prs.slides.add_slide(blank)
        rect(s, 0, 0, prs.slide_width, prs.slide_height, PAPER)
        rect(s, 0, 0, Inches(0.28), prs.slide_height, ACCENT)
        text(s, Inches(0.9), Inches(0.5), Inches(11.6), Inches(1.2), sl.title, 32 if len(sl.title) < 55 else 26, NAVY,
             bold=True, anchor=MSO_ANCHOR.MIDDLE)
        rect(s, Inches(0.9), Inches(1.75), Inches(1.2), Inches(0.06), ACCENT)
        if sl.bullets:
            n = len(sl.bullets)
            size = 26 if n <= 3 else 24 if n == 4 else 22
            box = s.shapes.add_textbox(Inches(0.9), Inches(2.1), Inches(11.4), Inches(4.5))
            tf = box.text_frame
            tf.word_wrap = True
            for j, b in enumerate(sl.bullets):
                p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                p.space_after = Pt(14)
                run = p.add_run()
                run.text = "•  " + b
                run.font.size, run.font.name = Pt(size), "Calibri"
                run.font.color.rgb = rgb(INK)
        text(s, Inches(11.6), Inches(6.85), Inches(1.4), Inches(0.4), f"{i} / {total}", 12, MUTED, align=PP_ALIGN.RIGHT)
        if sl.notes:
            s.notes_slide.notes_text_frame.text = sl.notes

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
