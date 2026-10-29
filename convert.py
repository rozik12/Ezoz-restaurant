"""File conversion between the formats students need. Content (headings, paragraphs, lists) is kept; fine layout is simplified."""
from __future__ import annotations

import io
import os
import re

FONT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "assets", "fonts")
TEXT_SOURCES = ("txt", "md", "docx", "pdf", "pptx")
IMAGE_SOURCES = ("jpg", "jpeg", "png", "webp", "bmp")
TARGETS = {                               # source extension -> offered targets
    "txt": ("docx", "pdf", "md", "pptx"),
    "md": ("docx", "pdf", "txt", "pptx"),
    "docx": ("pdf", "txt", "md", "pptx"),
    "pdf": ("docx", "txt", "md"),
    "pptx": ("docx", "pdf", "txt", "md"),
    **{e: ("pdf",) for e in IMAGE_SOURCES},
}
Block = tuple[str, str]                   # (kind, text); kind: h1 h2 h3 p li


class ConvertError(Exception):
    """`key` is an i18n key suffix (conv_err_<key>)."""

    def __init__(self, key: str) -> None:
        super().__init__(key)
        self.key = key


def ext_of(name: str) -> str:
    return name.rsplit(".", 1)[-1].lower() if "." in name else ""


# ---------------------------------------------------------------- reading
def _from_plain(text: str) -> list[Block]:
    blocks: list[Block] = []
    for chunk in re.split(r"\n\s*\n", text.replace("\r\n", "\n")):
        chunk = chunk.strip()
        if chunk:
            blocks.append(("p", re.sub(r"\s*\n\s*", " ", chunk)))
    return blocks


def _from_markdown(text: str) -> list[Block]:
    blocks: list[Block] = []
    para: list[str] = []

    def flush() -> None:
        if para:
            blocks.append(("p", " ".join(para)))
            para.clear()

    for line in text.replace("\r\n", "\n").split("\n"):
        s = line.strip()
        m = re.match(r"^(#{1,6})\s+(.*)$", s)
        li = re.match(r"^(?:[-*+•]|\d+[.)])\s+(.*)$", s)
        if m:
            flush()
            blocks.append((f"h{min(len(m.group(1)), 3)}", re.sub(r"[*_`]", "", m.group(2))))
        elif li:
            flush()
            blocks.append(("li", re.sub(r"[*_`]", "", li.group(1))))
        elif not s:
            flush()
        else:
            para.append(re.sub(r"[*_`]", "", s))
    flush()
    return blocks


def _from_docx(data: bytes) -> list[Block]:
    from docx import Document
    try:
        doc = Document(io.BytesIO(data))
    except Exception as e:
        raise ConvertError("read") from e
    blocks: list[Block] = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = (p.style.name or "").lower() if p.style is not None else ""
        if style == "title":
            blocks.append(("h1", text))
        elif style.startswith("heading"):
            n = "".join(c for c in style if c.isdigit()) or "1"
            blocks.append((f"h{min(int(n), 3)}", text))
        elif "list" in style:
            blocks.append(("li", text))
        else:
            blocks.append(("p", text))
    for table in doc.tables:                                   # tables become plain rows, in order after the text
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                blocks.append(("p", " | ".join(cells)))
    return blocks


def _from_pdf(data: bytes) -> list[Block]:
    from pypdf import PdfReader
    try:
        reader = PdfReader(io.BytesIO(data))
        text = "\n\n".join((page.extract_text() or "") for page in reader.pages)
    except Exception as e:
        raise ConvertError("read") from e
    if not text.strip():
        raise ConvertError("scan")
    return _from_plain(text)


def _from_pptx(data: bytes) -> list[Block]:
    from pptx import Presentation
    try:
        prs = Presentation(io.BytesIO(data))
    except Exception as e:
        raise ConvertError("read") from e
    blocks: list[Block] = []
    for slide in prs.slides:
        first = True
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                text = "".join(r.text for r in para.runs).strip()
                if text:
                    blocks.append(("h2", text) if first else ("li", text))
                    first = False
    return blocks


def read_blocks(name: str, data: bytes) -> list[Block]:
    ext = ext_of(name)
    if ext == "txt":
        blocks = _from_plain(data.decode("utf-8", errors="replace"))
    elif ext == "md":
        blocks = _from_markdown(data.decode("utf-8", errors="replace"))
    elif ext == "docx":
        blocks = _from_docx(data)
    elif ext == "pdf":
        blocks = _from_pdf(data)
    elif ext == "pptx":
        blocks = _from_pptx(data)
    else:
        raise ConvertError("unsupported")
    if not blocks:
        raise ConvertError("empty")
    return blocks


# ---------------------------------------------------------------- writing
def to_txt(blocks: list[Block]) -> bytes:
    out = []
    for kind, text in blocks:
        out.append(text.upper() if kind == "h1" else text if kind != "li" else f"• {text}")
    return "\n\n".join(out).encode("utf-8")


def to_md(blocks: list[Block]) -> bytes:
    out = []
    for kind, text in blocks:
        out.append({"h1": f"# {text}", "h2": f"## {text}", "h3": f"### {text}", "li": f"- {text}"}.get(kind, text))
    return "\n\n".join(out).encode("utf-8")


def to_docx(blocks: list[Block]) -> bytes:
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(14)
    for kind, text in blocks:
        if kind in ("h1", "h2", "h3"):
            doc.add_heading(text, level=int(kind[1]))
        elif kind == "li":
            doc.add_paragraph(text, style="List Bullet")
        else:
            doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def to_pdf(blocks: list[Block]) -> bytes:
    from fpdf import FPDF
    pdf = FPDF(format="A4")
    pdf.set_margins(20, 20, 15)
    pdf.set_auto_page_break(True, margin=18)
    pdf.add_font("DejaVu", "", os.path.join(FONT_DIR, "DejaVuSans.ttf"))
    pdf.add_font("DejaVu", "B", os.path.join(FONT_DIR, "DejaVuSans-Bold.ttf"))
    pdf.add_page()
    sizes = {"h1": 20, "h2": 16, "h3": 13}
    for kind, text in blocks:
        if kind in sizes:
            pdf.set_font("DejaVu", "B", sizes[kind])
            pdf.multi_cell(0, sizes[kind] * 0.55, text, new_x="LMARGIN", new_y="NEXT")
            pdf.ln(2)
        else:
            pdf.set_font("DejaVu", "", 11)
            if kind == "li":
                pdf.set_x(pdf.l_margin + 4)
                pdf.multi_cell(0, 6, "• " + text, new_x="LMARGIN", new_y="NEXT")
            else:
                pdf.multi_cell(0, 6, text, new_x="LMARGIN", new_y="NEXT")
            pdf.ln(2)
    return bytes(pdf.output())


def to_pptx(blocks: list[Block]) -> bytes:
    from slides import Deck, Slide, build_pptx
    title = next((t for k, t in blocks if k in ("h1", "h2")), "")
    slides: list[Slide] = []
    for kind, text in blocks:
        if kind in ("h1", "h2", "h3"):
            if kind == "h1" and text == title and not slides:
                continue
            slides.append(Slide(text, []))
        else:
            if not slides:
                slides.append(Slide(title or " ", []))
            if len(slides[-1].bullets) >= 5:
                slides.append(Slide(slides[-1].title, []))
            slides[-1].bullets.append(text if len(text) <= 160 else text[:157] + "…")
    return build_pptx(Deck(title or "Presentation", "", slides))


def image_to_pdf(data: bytes) -> bytes:
    from PIL import Image
    try:
        img = Image.open(io.BytesIO(data))
        img = img.convert("RGB")
    except Exception as e:
        raise ConvertError("read") from e
    buf = io.BytesIO()
    img.save(buf, format="PDF", resolution=150.0)
    return buf.getvalue()


WRITERS = {"txt": to_txt, "md": to_md, "docx": to_docx, "pdf": to_pdf, "pptx": to_pptx}


def convert(name: str, data: bytes, target: str) -> tuple[str, bytes]:
    """Returns (new file name, bytes)."""
    src = ext_of(name)
    if target not in TARGETS.get(src, ()):
        raise ConvertError("unsupported")
    base = name.rsplit(".", 1)[0] if "." in name else name
    if src in IMAGE_SOURCES:
        return f"{base}.pdf", image_to_pdf(data)
    blocks = read_blocks(name, data)
    return f"{base}.{target}", WRITERS[target](blocks)
